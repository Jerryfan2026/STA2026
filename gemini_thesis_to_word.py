"""
Google Gemini 调研报告生成器 - Word导出版本
支持直接生成符合中国大学格式的Word文档(.docx)
"""

import google.generativeai as genai
from typing import Optional, List, Dict
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
import matplotlib.pyplot as plt
import numpy as np
from datetime import datetime
import io
import os


class UniversityThesisWordGenerator:
    """生成符合中国大学格式的Word文档的论文生成器"""
    
    def __init__(self, api_key: str):
        """
        初始化生成器
        
        Args:
            api_key: Google Gemini API密钥
        """
        genai.configure(api_key=api_key)
        self.model = genai.GenerativeModel('gemini-pro')
        self.report_sections = {}
        self.charts = []
        self.references = []
        self.keywords = []
        self.doc = None
        
        # 设置中文字体
        plt.rcParams['font.sans-serif'] = ['SimHei', 'DejaVu Sans']
        plt.rcParams['axes.unicode_minus'] = False
    
    def _set_cell_border(self, cell, **kwargs):
        """设置表格单元格边框"""
        tcPr = cell._element.get_or_add_tcPr()
        tcBorders = OxmlElement('w:tcBorders')
        
        for edge in ('top', 'left', 'bottom', 'right'):
            if edge in kwargs:
                edge_data = kwargs.get(edge)
                edge_el = OxmlElement(f'w:{edge}')
                edge_el.set(qn('w:val'), 'single')
                edge_el.set(qn('w:sz'), '12')
                edge_el.set(qn('w:space'), '0')
                edge_el.set(qn('w:color'), '000000')
                tcBorders.append(edge_el)
        
        tcPr.append(tcBorders)
    
    def _add_heading(self, doc: Document, text: str, level: int):
        """添加标题"""
        heading = doc.add_heading(text, level=level)
        heading.paragraph_format.space_before = Pt(12)
        heading.paragraph_format.space_after = Pt(12)
        for run in heading.runs:
            run.font.name = '宋体'
            run.font.size = Pt(14 if level == 1 else 12)
    
    def _add_paragraph(self, doc: Document, text: str, indent: bool = False):
        """添加段落"""
        para = doc.add_paragraph(text)
        para.paragraph_format.line_spacing = 1.5
        para.paragraph_format.space_after = Pt(6)
        
        if indent:
            para.paragraph_format.first_line_indent = Inches(0.5)
        
        for run in para.runs:
            run.font.name = '宋体'
            run.font.size = Pt(12)
    
    def generate_chinese_abstract(self, topic: str, focus_areas: Optional[List[str]] = None) -> Dict[str, str]:
        """生成中文摘要"""
        focus_text = ""
        if focus_areas:
            focus_text = f"\n主要研究领域：{', '.join(focus_areas)}"
        
        prompt = f"""请为以下主题生成一份学位论文级别的中文摘要（200-300字）：

研究主题：{topic}{focus_text}

要求：
1. 简明扼要地说明研究背景和意义
2. 陈述研究主要内容和方法
3. 突出研究成果和创新之处
4. 指出研究的应用价值
5. 摘要语言严谨，符合学术规范

同时请提供5-8个关键词（用中文逗号分隔）。
回复格式：
【摘要】
<摘要内容>

【关键词】
<关键词列表>"""
        
        response = self.model.generate_content(prompt)
        text = response.text
        
        abstract = ""
        keywords = ""
        
        if "【摘要】" in text and "【关键词】" in text:
            abstract = text.split("【摘要】")[1].split("【关键词】")[0].strip()
            keywords = text.split("【关键词】")[1].strip()
        
        return {'abstract': abstract, 'keywords': keywords}
    
    def generate_english_abstract(self, topic: str, focus_areas: Optional[List[str]] = None) -> Dict[str, str]:
        """生成英文摘要"""
        focus_text = ""
        if focus_areas:
            focus_text = f"\nKey research areas: {', '.join(focus_areas)}"
        
        prompt = f"""Please generate an academic English abstract (200-250 words) for:

Topic: {topic}{focus_text}

Format:
【ABSTRACT】
<abstract content>

【KEYWORDS】
<keywords list>"""
        
        response = self.model.generate_content(prompt)
        text = response.text
        
        abstract = ""
        keywords = ""
        
        if "【ABSTRACT】" in text:
            abstract = text.split("【ABSTRACT】")[1].split("【KEYWORDS】")[0].strip() if "【KEYWORDS】" in text else text.split("【ABSTRACT】")[1].strip()
            keywords = text.split("【KEYWORDS】")[1].strip() if "【KEYWORDS】" in text else ""
        
        return {'abstract': abstract, 'keywords': keywords}
    
    def generate_introduction(self, topic: str) -> str:
        """生成引言"""
        prompt = f"""请为以下研究主题生成符合学位论文规范的引言部分（500-800字）：

研究主题：{topic}

要求：
1. 介绍研究课题的理论背景
2. 分析问题的现实意义
3. 阐述国内外研究现状
4. 明确提出研究目标
5. 概括论文主要内容"""
        
        response = self.model.generate_content(prompt)
        return response.text
    
    def generate_literature_review(self, topic: str, key_concepts: Optional[List[str]] = None) -> str:
        """生成文献综述"""
        concept_text = ""
        if key_concepts:
            concept_text = f"\n主要概念：{', '.join(key_concepts)}"
        
        prompt = f"""请为以下研究主题生成学位论文级别的文献综述（1000-1500字）：

研究主题：{topic}{concept_text}

要求：
1. 系统梳理国内外相关研究现状
2. 分析不同学者的主要观点
3. 指出当前研究的优势和不足
4. 明确阐述理论基础
5. 结合具体文献进行论述"""
        
        response = self.model.generate_content(prompt)
        return response.text
    
    def generate_methodology(self, topic: str, research_methods: List[str]) -> str:
        """生成研究方法"""
        methods_text = f"研究方法：{', '.join(research_methods)}"
        
        prompt = f"""请为以下研究生成符合学位论文规范的研究方法部分（400-600字）：

研究主题：{topic}
{methods_text}

要求：
1. 明确阐述采用的研究方法和工具
2. 说明数据来源和收集方式
3. 描述数据处理和分析过程
4. 阐述研究的科学性和可靠性"""
        
        response = self.model.generate_content(prompt)
        return response.text
    
    def generate_analysis(self, topic: str, focus_areas: Optional[List[str]] = None) -> str:
        """生成分析部分"""
        focus_text = ""
        if focus_areas:
            focus_text = f"\n分析重点：{', '.join(focus_areas)}"
        
        prompt = f"""请为以下研究��题生成学位论文级别的分析部分（1500-2000字）：

研究主题：{topic}{focus_text}

要求：
1. 多维度、深层次地分析研究问题
2. 运用数据、案例进行论证
3. 逐步深化分析层次
4. 提出具有见地的观点
5. 分析过程严密，论证充分"""
        
        response = self.model.generate_content(prompt)
        return response.text
    
    def generate_conclusions(self, topic: str, findings: Optional[List[str]] = None) -> str:
        """生成结论"""
        findings_text = ""
        if findings:
            findings_text = f"\n主要发现：\n" + "\n".join([f"- {finding}" for finding in findings])
        
        prompt = f"""请为以下研究生成学位论文级别的结论与建议部分（400-600字）：

研究主题：{topic}{findings_text}

要求：
1. 准确概括研究的主要结论
2. 阐述结论的理论和实践意义
3. 说明研究的局限性
4. 针对问题提出切实可行的建议
5. 指出未来的研究方向"""
        
        response = self.model.generate_content(prompt)
        return response.text
    
    def generate_references(self, topic: str, count: int = 20) -> List[str]:
        """生成参考文献"""
        prompt = f"""请为以下研究主题生成{count}条符合GB/T 7714-2015标准的参考文献：

研究主题：{topic}

要求：
1. 包含中文期刊论文、中文著作、外文论文等多种类型
2. 遵循GB/T 7714-2015参考文献著录规则
3. 按照在论文中出现的顺序编号

请直接返回按[1][2]...[n]编号的参考文献列表"""
        
        response = self.model.generate_content(prompt)
        ref_text = response.text
        
        references = []
        lines = ref_text.strip().split('\n')
        for line in lines:
            if line.strip() and '[' in line:
                references.append(line.strip())
        
        return references
    
    def create_sample_charts(self) -> List[str]:
        """创建示例图表"""
        chart_files = []
        
        # 图表1：趋势图
        fig, ax = plt.subplots(figsize=(10, 6))
        fig.patch.set_facecolor('white')
        ax.set_facecolor('#f5f5f5')
        
        years = [2019, 2020, 2021, 2022, 2023]
        values = [15, 28, 42, 65, 88]
        
        ax.plot(years, values, marker='o', linewidth=2.5, markersize=8, color='#2E86AB')
        for year, value in zip(years, values):
            ax.text(year, value + 2, f'{value}%', ha='center', va='bottom', fontsize=10)
        
        ax.set_xlabel('年份', fontsize=11, fontweight='bold')
        ax.set_ylabel('增长率 (%)', fontsize=11, fontweight='bold')
        ax.set_title('图1 近五年发展趋势分析', fontsize=12, fontweight='bold')
        ax.grid(True, alpha=0.3)
        ax.set_ylim(0, max(values) + 15)
        
        filename = 'fig_1_trend.png'
        plt.savefig(filename, dpi=300, bbox_inches='tight', facecolor='white')
        plt.close()
        chart_files.append(filename)
        
        # 图表2：柱状图
        fig, ax = plt.subplots(figsize=(10, 6))
        fig.patch.set_facecolor('white')
        ax.set_facecolor('#f5f5f5')
        
        categories = ['地区A', '地区B', '地区C', '地区D']
        values = [28, 35, 22, 15]
        colors = ['#A23B72', '#F18F01', '#C73E1D', '#6A994E']
        
        bars = ax.bar(categories, values, color=colors, edgecolor='black', linewidth=1.2, alpha=0.85)
        for bar, value in zip(bars, values):
            height = bar.get_height()
            ax.text(bar.get_x() + bar.get_width()/2., height,
                   f'{value}%', ha='center', va='bottom', fontsize=10, fontweight='bold')
        
        ax.set_ylabel('占比 (%)', fontsize=11, fontweight='bold')
        ax.set_title('图2 不同地区分布情况', fontsize=12, fontweight='bold')
        ax.grid(True, alpha=0.3, axis='y')
        ax.set_ylim(0, max(values) + 10)
        
        filename = 'fig_2_distribution.png'
        plt.savefig(filename, dpi=300, bbox_inches='tight', facecolor='white')
        plt.close()
        chart_files.append(filename)
        
        # 图表3：饼图
        fig, ax = plt.subplots(figsize=(10, 8))
        fig.patch.set_facecolor('white')
        
        labels = ['研发投入', '生产成本', '销售渠道', '服务支持']
        sizes = [25, 30, 25, 20]
        colors = ['#FF6B6B', '#4ECDC4', '#45B7D1', '#FFA07A']
        
        ax.pie(sizes, labels=labels, autopct='%1.1f%%', colors=colors, startangle=90)
        ax.set_title('图3 产业链构成分析', fontsize=12, fontweight='bold')
        
        filename = 'fig_3_composition.png'
        plt.savefig(filename, dpi=300, bbox_inches='tight', facecolor='white')
        plt.close()
        chart_files.append(filename)
        
        self.charts = chart_files
        return chart_files
    
    def generate_complete_thesis(self,
                                title: str,
                                author: str,
                                school: str,
                                major: str,
                                advisor: str,
                                topic: str,
                                focus_areas: Optional[List[str]] = None,
                                research_methods: Optional[List[str]] = None,
                                key_findings: Optional[List[str]] = None) -> Document:
        """
        生成完整的学位论文Word文档
        
        Args:
            title: 论文标题
            author: 作者名字
            school: 学校名称
            major: 专业/学院
            advisor: 指导老师
            topic: 研究主题
            focus_areas: 研究重点领域
            research_methods: 研究方法
            key_findings: 关键发现
            
        Returns:
            Word文档对象
        """
        if research_methods is None:
            research_methods = ['文献分析法', '数据统计法', '对比分析法']
        
        print(f"🎓 开始生成『{title}』的学位论文...\n")
        
        # 创建Word文档
        self.doc = Document()
        
        # 设置默认字体
        style = self.doc.styles['Normal']
        style.font.name = '宋体'
        style.font.size = Pt(12)
        
        # 生成封面
        print("📖 生成封面...")
        self._create_cover(title, author, school, major, advisor)
        self.doc.add_page_break()
        
        # 生成中文摘要
        print("📝 生成中文摘要...")
        abstract_dict = self.generate_chinese_abstract(topic, focus_areas)
        self._add_heading(self.doc, '中文摘要', 1)
        self._add_paragraph(self.doc, abstract_dict['abstract'], indent=True)
        self._add_paragraph(self.doc, f"关键词：{abstract_dict['keywords']}")
        self.doc.add_page_break()
        
        # 生成英文摘要
        print("📝 生成英文摘要...")
        en_abstract_dict = self.generate_english_abstract(topic, focus_areas)
        self._add_heading(self.doc, 'ABSTRACT', 1)
        self._add_paragraph(self.doc, en_abstract_dict['abstract'], indent=True)
        self._add_paragraph(self.doc, f"Keywords: {en_abstract_dict['keywords']}")
        self.doc.add_page_break()
        
        # 生成目录
        print("📋 生成目录...")
        self._add_heading(self.doc, '目    录', 1)
        toc_items = [
            '1. 引言',
            '2. 文献综述',
            '3. 研究方法',
            '4. 分析讨论',
            '5. 结论与建议',
            '参考文献',
            '附录'
        ]
        for item in toc_items:
            self._add_paragraph(self.doc, item)
        self.doc.add_page_break()
        
        # 生成引言
        print("📚 生成引言...")
        introduction = self.generate_introduction(topic)
        self._add_heading(self.doc, '1. 引    言', 1)
        self._add_paragraph(self.doc, introduction, indent=True)
        
        # 生成文献综述
        print("📚 生成文献综述...")
        literature_review = self.generate_literature_review(topic, focus_areas)
        self._add_heading(self.doc, '2. 文 献 综 述', 1)
        self._add_paragraph(self.doc, literature_review, indent=True)
        
        # 生成研究方法
        print("🔬 生成研究方法...")
        methodology = self.generate_methodology(topic, research_methods)
        self._add_heading(self.doc, '3. 研 究 方 法', 1)
        self._add_paragraph(self.doc, methodology, indent=True)
        
        # 生成分析
        print("🔍 生成分析部分...")
        analysis = self.generate_analysis(topic, focus_areas)
        self._add_heading(self.doc, '4. 分 析 讨 论', 1)
        self._add_paragraph(self.doc, analysis, indent=True)
        
        # 生成结论
        print("💡 生成结论与建议...")
        conclusions = self.generate_conclusions(topic, key_findings)
        self._add_heading(self.doc, '5. 结 论 与 建 议', 1)
        self._add_paragraph(self.doc, conclusions, indent=True)
        
        # 生成参考文献
        print("📋 生成参考文献...")
        references = self.generate_references(topic, count=20)
        self._add_heading(self.doc, '参 考 文 献', 1)
        for ref in references:
            self._add_paragraph(self.doc, ref)
        
        # 生成图表
        print("📊 生成数据图表...\n")
        self._create_sample_charts()
        self._add_heading(self.doc, '附    录', 1)
        self._add_heading(self.doc, '附录 A：图表', 2)
        for idx, chart_file in enumerate(self.charts, 1):
            try:
                self._add_paragraph(self.doc, f"图 {idx}")
                self.doc.add_picture(chart_file, width=Inches(5.5))
                self._add_paragraph(self.doc, "")
            except Exception as e:
                print(f"无法插入图表 {chart_file}: {e}")
        
        # 添加致谢
        self.doc.add_page_break()
        self._add_heading(self.doc, '致    谢', 1)
        thanks_text = f"""
本研究的完成离不开各位指导老师的耐心指导和同学的大力支持。
特别感谢 {advisor} 老师的精心指导，以及 {school} 提供的良好研究环境。
同时感谢所有参与问卷调查和访谈的受访者。
"""
        self._add_paragraph(self.doc, thanks_text)
        
        print("✓ 论文生成完成\n")
        return self.doc
    
    def _create_cover(self, title: str, author: str, school: str, major: str, advisor: str):
        """创建论文封面"""
        # 添加学校名称
        cover_title = self.doc.add_paragraph()
        cover_title.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
        cover_title.paragraph_format.space_before = Pt(40)
        run = cover_title.add_run(school)
        run.font.size = Pt(18)
        run.font.bold = True
        run.font.name = '宋体'
        
        # 添加论文类型
        thesis_type = self.doc.add_paragraph()
        thesis_type.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
        thesis_type.paragraph_format.space_before = Pt(60)
        run = thesis_type.add_run('调研报告（学位论文）')
        run.font.size = Pt(16)
        run.font.name = '宋体'
        
        # 添加论文标题
        thesis_title = self.doc.add_paragraph()
        thesis_title.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
        thesis_title.paragraph_format.space_before = Pt(60)
        run = thesis_title.add_run(title)
        run.font.size = Pt(14)
        run.font.bold = True
        run.font.name = '宋体'
        
        # 添加作者信息表格
        info_para = self.doc.add_paragraph()
        info_para.paragraph_format.space_before = Pt(60)
        
        table = self.doc.add_table(rows=4, cols=2)
        table.style = 'Light Grid Accent 1'
        
        # 设置表格内容
        table.cell(0, 0).text = '学院/系部'
        table.cell(0, 1).text = major
        table.cell(1, 0).text = '作者姓名'
        table.cell(1, 1).text = author
        table.cell(2, 0).text = '指导教师'
        table.cell(2, 1).text = advisor
        table.cell(3, 0).text = '完成日期'
        table.cell(3, 1).text = datetime.now().strftime('%Y年%m月%d日')
        
        # 设置表格格式
        for row in table.rows:
            for cell in row.cells:
                for paragraph in cell.paragraphs:
                    paragraph.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
                    for run in paragraph.runs:
                        run.font.name = '宋体'
                        run.font.size = Pt(11)
    
    def _create_sample_charts(self):
        """创建示例图表（包装方法）"""
        self.create_sample_charts()
    
    def save_to_word(self, filename: str = "学位论文.docx"):
        """
        保存为Word文档
        
        Args:
            filename: 输出文件名
        """
        if self.doc is None:
            print("❌ 错误：尚未生成文档，请先调用 generate_complete_thesis()")
            return
        
        try:
            self.doc.save(filename)
            print(f"✅ Word文档已保存到: {filename}")
        except Exception as e:
            print(f"❌ 保存文档时出错: {e}")


# 使用示例
def main():
    """主函数"""
    
    API_KEY = "YOUR_GEMINI_API_KEY"  # 替换为你的API密钥
    
    generator = UniversityThesisWordGenerator(api_key=API_KEY)
    
    # 生成论文
    thesis = generator.generate_complete_thesis(
        title="人工智能在医疗健康领域应用的现状与展望",
        author="张三",
        school="北京大学",
        major="计算机学院 / 人工智能专业",
        advisor="李四教授",
        topic="人工智能在医疗健康中的应用",
        focus_areas=[
            "AI诊断技术",
            "个性化治疗方案",
            "医疗成本优化",
            "数据隐私保护"
        ],
        research_methods=[
            "文献分析法",
            "案例研究法",
            "数据统计法",
            "对比分析法"
        ],
        key_findings=[
            "AI诊断准确率已达到95%以上",
            "个性化治疗将成为未来医疗主流方向",
            "需要建立完善的数据安全保护机制",
            "应用前景广阔但面临伦理挑战"
        ]
    )
    
    # 保存为Word
    generator.save_to_word("学位论文_标准格式.docx")
    print("\n✨ 报告生成完毕！")


if __name__ == "__main__":
    main()
